# Sydney Fuller one-page contractor offer (terms only) built from the DWZ proposal's own .docx export: every
# paragraph/table is a deep copy of a DWZ element, only the text changes.
import copy, docx
from docx.text.paragraph import Paragraph
from docx.table import Table
d = docx.Document('dwz.docx')
body = d.element.body
E = list(body.iterchildren())
sect = E[-1]
out = []

def settext(p, texts):
    runs = p.runs
    for k, r in enumerate(runs):
        r.text = texts[k] if k < len(texts) else ''

def p(i, *texts):
    el = copy.deepcopy(E[i]); settext(Paragraph(el, d), texts); out.append(el)

def bullets(items):
    for it in items: p(17, it)

def table(i, rows, bold_last=False):
    el = copy.deepcopy(E[i]); t = Table(el, d)
    protos = [copy.deepcopy(r._tr) for r in t.rows]
    for r in list(t.rows): el.remove(r._tr)
    for n, row in enumerate(rows):
        tr = copy.deepcopy(protos[0] if n == 0 else (protos[-1] if bold_last and n == len(rows) - 1 else protos[1]))
        el.append(tr)
    t = Table(el, d)
    for r, row in zip(t.rows, rows):
        for c, txt in zip(r.cells, row):
            settext(c.paragraphs[0], [txt])
    out.append(el)
p(0, 'CONTRACTOR OFFER · SUMMARY OF TERMS')
p(1, 'Sydney Fuller')
p(2, 'Producing Consultant / Jr. Producer — Straw Hut Media')
p(4, 'Prepared by: ', 'Ryan Tillotson, Founder & CEO, Straw Hut Media')
p(5, 'Date: ', 'September 25, 2026')
p(9, 'The Terms')
table(37, [['Term', 'Detail'],
    ['Role', 'Producing Consultant / Jr. Producer, focused on CodeStrap (Code x Connor), plus other projects as offered'],
    ['Trial period', 'Three months, October 1 – December 31, 2026'],
    ['Rate', '$20 per hour'],
    ['Monthly minimum', '$1,000 per month, even if fewer hours are needed'],
    ['Monthly maximum', '$1,750 per month (87.5 hours); more only with Ryan’s written approval'],
    ['Paid', 'Within 10 days after each month ends, based on the hours in your tracking spreadsheet (no invoice needed)']])
p(28, 'What You Need to Do')
bullets(['Track your hours and tasks as you work in the shared Straw Hut tracking spreadsheet',
    'Keep the spreadsheet up to date by the last day of each month; we pay from it',
    'Send a completed Form W-9 before your first payment',
    'Keep all client and show information confidential'])
p(28, 'What “Independent Contractor” Means')
bullets(['You are not an employee: no taxes are withheld, and you pay your own income and self-employment taxes (1099-NEC)',
    'No employee benefits: no health insurance, paid time off, sick leave, overtime, workers’ comp, or unemployment insurance',
    'You set your own hours and methods around recording dates and deadlines, and may work for other clients',
    'Either of us can end the arrangement with seven days’ written notice'])
p(28, 'Next Step')
p(16, 'If these terms work for you, sign below. We’ll then send the full Independent Contractor Agreement, covering these terms plus confidentiality and ownership of work.')
p(69, 'Agreed — Sydney Fuller')
p(70, 'Signature: ______________________________    Date: ______________')

for el in E[:-1]: body.remove(el)
for el in out: body.insert(len(body) - 1, el)
for fp in d.sections[0].footer.paragraphs:
    for r in fp.runs:
        if 'DIE WITH ZERO' in r.text:
            r.text = r.text.replace('DIE WITH ZERO PODCAST', 'SYDNEY FULLER').replace('STRAW HUT MEDIA PROPOSAL', 'CONTRACTOR OFFER')
d.core_properties.title = 'Sydney Fuller Contractor Offer'
d.save('syd_terms.docx')
