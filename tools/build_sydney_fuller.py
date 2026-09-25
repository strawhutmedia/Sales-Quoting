# Sydney Fuller contractor agreement built from the DWZ proposal's own .docx export: every
# paragraph/table is a deep copy of a DWZ element, only the text changes.
import copy, docx
from docx.text.paragraph import Paragraph
from docx.table import Table
d = docx.Document('dwz.docx')
body = d.element.body
E = list(body.iterchildren())
sect = E[-1]
out = []

def settext(p, texts):
    runs = p.runs
    for k, r in enumerate(runs):
        r.text = texts[k] if k < len(texts) else ''

def p(i, *texts):
    el = copy.deepcopy(E[i]); settext(Paragraph(el, d), texts); out.append(el)

def bullets(items):
    for it in items: p(17, it)

def table(i, rows, bold_last=False):
    el = copy.deepcopy(E[i]); t = Table(el, d)
    protos = [copy.deepcopy(r._tr) for r in t.rows]
    for r in list(t.rows): el.remove(r._tr)
    for n, row in enumerate(rows):
        tr = copy.deepcopy(protos[0] if n == 0 else (protos[-1] if bold_last and n == len(rows) - 1 else protos[1]))
        el.append(tr)
    t = Table(el, d)
    for r, row in zip(t.rows, rows):
        for c, txt in zip(r.cells, row):
            settext(c.paragraphs[0], [txt])
    out.append(el)



p(0, 'INDEPENDENT CONTRACTOR AGREEMENT')
p(1, 'Sydney Fuller')
p(2, 'Contract Producer — Straw Hut Media')
p(3, 'Prepared for: ', 'Sydney Fuller')
p(4, 'Prepared by: ', 'Ryan Tillotson, Founder & CEO, Straw Hut Media')
p(5, 'Date: ', 'September 25, 2026')
p(6)
p(7, 'This Independent Contractor Agreement engages Sydney Fuller (“Contractor”) as an independent contractor to provide podcast producing services to Straw Hut Media for a three-month trial period, effective October 1, 2026.')
p(8); p(9, 'Engagement Details')
bullets(['Role: Contract Producer, primarily producing CodeStrap episodes, plus other production projects Straw Hut Media may offer',
    'Term: three-month trial, October 1, 2026 through December 31, 2026',
    'Fee: $1,733 per month, flat',
    'Expected commitment: approximately 20 hours per week on average',
    'After the trial: the engagement may be extended or revised by written agreement'])
p(28, 'Services')
p(16, 'Contractor will provide producing services for the podcast productions Straw Hut Media assigns, primarily CodeStrap, and may take on other production projects as offered. Contractor decides the methods, tools, and schedule used to complete the work, working around client recording dates and delivery deadlines. Contractor will communicate with the Straw Hut Media team primarily through Slack, stay reasonably responsive during active projects, and flag early if a deadline cannot be met.')
p(28, 'Payment')
p(16, 'Straw Hut Media will pay $1,733 for each month of services within 10 days of receiving Contractor’s invoice at the end of that month. Partial months are prorated. No expenses are reimbursed unless pre-approved in writing. Contractor will provide a completed Form W-9 before the first payment and will receive a Form 1099-NEC for the year.')
p(28, 'Independent Contractor Relationship')
p(16, 'Contractor is an independent contractor and not an employee, partner, or agent of Straw Hut Media, and has no authority to bind Straw Hut Media in any manner. Contractor understands and agrees that this means:')
bullets(['No taxes are withheld from payments; Contractor is solely responsible for all income and self-employment taxes, withholdings, and other statutory obligations',
    'Contractor is not eligible for employee benefits, including health insurance, paid time off, sick leave, overtime, workers’ compensation, or unemployment insurance through Straw Hut Media',
    'Contractor is responsible for her own insurance and for the equipment she needs to do the work (computer, phone, software); studio equipment is available for sessions',
    'Contractor sets her own working hours and methods, apart from client recording dates and deadlines',
    'Contractor is free to work for other clients, as long as it does not conflict with the confidentiality terms below',
    'Contractor will not represent herself as an employee of Straw Hut Media'])
p(28, 'Professional Conduct')
p(16, 'Contractor will represent Straw Hut Media in client-facing work, including with CodeStrap, its hosts, and guests, and will uphold Straw Hut Media’s standards of professionalism, creativity, and discretion.')
p(28, 'Confidentiality')
p(16, 'Contractor will keep confidential all non-public information acquired through this work, including unreleased episodes, raw audio and video, client and guest information, budgets, release dates, passwords, and internal processes, and will use it only to perform this work. Contractor will not post or share behind-the-scenes information or project content, including screenshots or vague references, on social media or elsewhere without written approval. Contractor will not use this information to solicit Straw Hut Media’s clients. These obligations survive the end of this agreement. Unauthorized disclosure may cause irreparable harm, and Straw Hut Media may seek injunctive relief.')
p(28, 'Ownership of Work')
p(16, 'All materials and work product Contractor creates or contributes under this agreement are “work made for hire” and belong exclusively to Straw Hut Media (or its clients). To the extent any work does not qualify as work made for hire, Contractor assigns all right, title, and interest in it to Straw Hut Media. Contractor may use released work in her portfolio only with Straw Hut Media’s written approval.')
p(28, 'Ending the Agreement')
p(16, 'Either party may end this agreement at any time with seven (7) days’ written notice. Straw Hut Media will pay for services performed through the end date, prorated. On ending, Contractor will deliver all project files and return any Straw Hut Media equipment, drives, or materials. This agreement is governed by California law, is the entire agreement between the parties, and may be changed only in writing signed by both parties.')
p(59, 'Contractor Acknowledgment')
p(60, 'By signing below, Contractor confirms that she has read and understands this agreement and specifically understands that:')
bullets(['this is an independent contractor role, not employment, for a three-month trial at a flat $1,733 per month',
    'no taxes will be withheld, and she is responsible for paying her own income and self-employment taxes',
    'she will not receive employee benefits, overtime, workers’ compensation, or unemployment insurance',
    'either party may end the engagement with seven days’ written notice',
    'she has had the opportunity to ask questions and to consult an advisor of her choosing before signing'])
p(61)
for i, t in zip(range(62, 67), ['Straw Hut Media', 'By: ______________________________________', 'Name: Ryan Tillotson', 'Title: Founder & CEO', 'Date: __________________']): p(i, t)
p(67)
for i, t in zip(range(69, 74), ['Contractor', 'Signature: ______________________________________', 'Name: Sydney Fuller', 'Title: Independent Contractor', 'Date: __________________']): p(i, t)
p(75, 'Ryan Tillotson')
p(76, 'Founder & CEO, Straw Hut Media')
p(77, 'ryan@strawhutmedia.com  ·  760.807.7138  ·  7201 Melrose Ave., Suite 203, Los Angeles, CA 90046')

for el in E[:-1]: body.remove(el)
for el in out: body.insert(len(body) - 1, el)
for fp in d.sections[0].footer.paragraphs:
    for r in fp.runs:
        if 'DIE WITH ZERO' in r.text:
            r.text = r.text.replace('DIE WITH ZERO PODCAST', 'SYDNEY FULLER').replace('STRAW HUT MEDIA PROPOSAL', 'CONTRACTOR AGREEMENT')
d.core_properties.title = 'Sydney Fuller Contractor Agreement'
d.save('sydney.docx')
