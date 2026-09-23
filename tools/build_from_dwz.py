# Flying V proposal built from the DWZ proposal's own .docx export: every
# paragraph/table is a deep copy of a DWZ element, only the text changes.
import copy, docx
from docx.text.paragraph import Paragraph
from docx.table import Table
d = docx.Document('dwz.docx')
body = d.element.body
E = list(body.iterchildren())
sect = E[-1]
out = []

def settext(p, texts):
    runs = p.runs
    for k, r in enumerate(runs):
        r.text = texts[k] if k < len(texts) else ''

def p(i, *texts):
    el = copy.deepcopy(E[i]); settext(Paragraph(el, d), texts); out.append(el)

def bullets(items):
    for it in items: p(17, it)

def table(i, rows, bold_last=False):
    el = copy.deepcopy(E[i]); t = Table(el, d)
    protos = [copy.deepcopy(r._tr) for r in t.rows]
    for r in list(t.rows): el.remove(r._tr)
    for n, row in enumerate(rows):
        tr = copy.deepcopy(protos[0] if n == 0 else (protos[-1] if bold_last and n == len(rows) - 1 else protos[1]))
        el.append(tr)
    t = Table(el, d)
    for r, row in zip(t.rows, rows):
        for c, txt in zip(r.cells, row):
            settext(c.paragraphs[0], [txt])
    out.append(el)

p(0, 'A FULL-SERVICE PODCAST PARTNERSHIP')
p(1, 'Flying V Group')
p(2, 'Full-Service Podcast Production Partnership — Straw Hut Media')
p(3, 'Prepared for: ', 'Flying V Group / Robb Fahrion')
p(4, 'Prepared by: ', 'Ryan Tillotson, Founder & CEO, Straw Hut Media')
p(5, 'Date: ', 'September 23, 2026')
p(6)
p(7, 'A single, full-service production partnership — editing, branding, theme music, motion graphics, social, audio versioning, distribution, and ongoing channel management across social, RSS, and YouTube — built so Flying V Group records the conversation and Straw Hut Media handles everything after it.')
p(8); p(9, 'Program Details')
bullets(['Show title: to be developed together during pre-production',
    'Host: Robb Fahrion, Co-Founder & CEO, Flying V Group',
    'Format: audio and video',
    'Recording: captured in-house at the Flying V Group offices, media sent to Straw Hut Media',
    'Release cadence: up to two (2) episodes per month',
    'Production begins: November 1, 2026 (production build-out and branding)',
    'Launch date: January 4, 2027',
    'Term: 12-month minimum agreement'])
p(14); p(15, 'Startup & Transition')
p(16, 'Flying V Group shoots and releases the first three (3) episodes of the show on its own. Straw Hut Media takes over production from there, beginning with the fourth episode and continuing through the remainder of the term.')
p(16, 'The period from November 1, 2026 to the January 4, 2027 launch serves as a startup and build period, during which Straw Hut Media focuses on show development and title, branding and cover art, original theme music, the motion graphics package, the capture spec for the Flying V Group offices, channel and feed setup, and launch preparation — so that when we assume production, the look, the sound, the workflow, and the distribution are already built and running. The monthly retainer applies from November 1, including the startup period before Straw Hut Media begins producing episodes.')
p(28, "What's Included")
p(16, "Straw Hut Media's full-service production partnership is delivered under a flat monthly retainer, described in full under Investment & Costs at the end of this proposal. The retainer covers:")
bullets(['Full episode editing: audio and video, multi-track editing, sound design, cleaned and mastered',
    'Original theme music composed for the show',
    'Custom intros, outros, and transitions built specifically for the show',
    'Branding: cover art, logo treatment, episode graphics, and social graphics',
    'Motion graphics: animated lower thirds, titles, and end cards',
    'Social clip selection, captioning, and copywriting for native, ready-to-share assets',
    'Audio versioning of each episode for audio-only release',
    'Distribution support and release scheduling to all major platforms',
    "Channel management: ongoing management of the show's social media channels, RSS feed, and YouTube channel, including publishing, scheduling, and channel upkeep",
    'Analytics to track download and audience growth'])
p(27, 'Cadence & Retainer Adjustment: This retainer is structured around a schedule of up to two episodes per month for the full 12-month term. Should release frequency increase beyond two episodes per month at any point during the term, the monthly retainer will be renegotiated prior to the new cadence taking effect.')
p(28, 'Your Team')
p(29, 'The monthly retainer covers dedicated time from the following Straw Hut Media roles for the duration of the engagement:')
bullets(['Project Coordinator', 'Producer', 'Editor', 'Motion Designer'])
p(35, 'Investment & Costs')
p(36, "This partnership runs on a single flat monthly retainer, which is Straw Hut Media's fee. Because Flying V Group records in-house at its own offices, there are no separate pass-through production costs — no studio rental, crew, lighting, or travel — attached to this engagement.")
table(37, [
    ['Cost Component', 'Details'],
    ['Monthly Retainer', '$2,450 / month, billed monthly beginning November 1, 2026. Covers full episode editing, original theme music, custom intros and outros, branding (cover art, logo treatment, episode and social graphics), motion graphics, social clip creation and copywriting, audio versioning, channel management (social, RSS, YouTube), distribution support, and analytics — for up to two episodes per month across the full 12-month term.'],
    ['Production Costs', 'None. Recording is captured in-house by Flying V Group at its own offices, so there are no studio rental, crew, lighting, or travel costs under this engagement.'],
    ['Optional — Studio Upgrade', '$4,350 / month in place of the standard retainer, should Flying V Group prefer to record at the Straw Hut Media studio in Hollywood. See the Recording Avenues section below.']])
p(38)
p(39, 'Monthly Retainer: ', '$2,450 / month', '  ·  billed monthly, 12-month minimum term')
p(16, 'Prior to the production start date, Flying V Group provides a credit card or bank account (ACH) on file, and authorizes Straw Hut Media to charge that method for the monthly retainer at the beginning of each month of the term.')
p(40, 'Recording Avenues')
p(41, 'The show can be captured one of two ways. Both deliver the same finished product — the same editing, theme music, branding, motion graphics, social assets, and channel management — and differ only in where and how the recording happens.')
p(42)
p(44, 'In-House Avenue')
p(45, "Flying V Group records in its own offices, on its own schedule, and handles its own recording sessions, guest booking, and scheduling. Straw Hut Media provides a simple capture spec up front — camera framing, audio settings, and file format — so the media arrives clean and consistent every time. Files are dropped into a shared folder, and that is the end of Flying V Group's lift for the episode. This is the recommended starting point: it is the lowest cost, the fastest to stand up, and it keeps recording entirely on your calendar rather than ours.")
p(46, 'Studio Avenue')
p(47, 'Flying V Group records at the Straw Hut Media studio in Hollywood. In addition to everything in the In-House Avenue, the studio package includes multi-camera video recording in our space, guest booking and scheduling, a dedicated production manager to run each shoot, and priority support. This avenue is worth considering if the show leans heavily on guests, or if a consistent, controlled studio look matters to how the brand presents.')
p(48)
p(49, 'Avenue Comparison')
table(51, [
    ['', 'In-House Avenue', 'Studio Avenue'],
    ['Where you record', 'Flying V Group offices', 'Straw Hut studio, Hollywood'],
    ['Cameras & crew', 'Your setup, to our capture spec', 'Multi-camera, our crew'],
    ['Editing, branding & motion graphics', 'Included', 'Included'],
    ['Original theme music', 'Included', 'Included'],
    ['Social clips & channel management', 'Included', 'Included'],
    ['Guest booking & scheduling', 'Flying V Group', 'Included'],
    ['Dedicated production manager', '—', 'Included'],
    ['Monthly retainer', '$2,450', '$4,350']], bold_last=True)
p(52)
p(53, 'Flying V Group can move between avenues during the term with written agreement — starting in-house and stepping up to the studio later does not require restarting the engagement.')
p(57, 'Next Steps')
p(58, "We're glad to walk through the format in more detail, lock the capture spec for your offices, and set a first recording date. Let us know which avenue makes the most sense as we move toward the November 1 production start, and we'll send a short services agreement reflecting these terms and get branding underway.")
p(59, 'Acceptance')
p(60, 'By signing below, both parties agree to move forward on the terms outlined in this proposal. A formal, long-form agreement reflecting these terms will be drafted and executed prior to the start of production.')
p(61)
for i, t in zip(range(62, 67), ['Straw Hut Media', 'By: ______________________________________', 'Name: Ryan Tillotson', 'Title: Founder & CEO', 'Date: __________________']): p(i, t)
p(67); p(68)
for i, t in zip(range(69, 74), ['Flying V Group', 'By: ______________________________________', 'Name: Robb Fahrion', 'Title: Co-Founder & CEO', 'Date: __________________']): p(i, t)
p(74)
p(75, 'Ryan Tillotson')
p(76, 'Founder & CEO, Straw Hut Media')
p(77, 'ryan@strawhutmedia.com  ·  760.807.7138  ·  7201 Melrose Ave., Suite 203, Los Angeles, CA 90046')

for el in E[:-1]: body.remove(el)
for el in out: body.insert(len(body) - 1, el)
for fp in d.sections[0].footer.paragraphs:
    for r in fp.runs:
        if 'DIE WITH ZERO' in r.text: r.text = r.text.replace('DIE WITH ZERO PODCAST', 'FLYING V GROUP PODCAST')
d.core_properties.title = 'Flying V Group Podcast Proposal'
d.save('fv.docx')
t = '\n'.join(x.text for x in docx.Document('fv.docx').paragraphs)
print('saved', 'Die With Zero' in t, 'Dillon' in t, [f.text for f in docx.Document('fv.docx').sections[0].footer.paragraphs])
